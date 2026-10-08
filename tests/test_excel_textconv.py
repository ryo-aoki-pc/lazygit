"""Exercise the public converter and the Git workflows used by lazygit.

Run with: python3 -m unittest discover -s lazygit/tests -v
"""

import datetime as dt
import hashlib
import io
import os
from pathlib import Path
import shlex
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


CONVERTER = Path(__file__).resolve().parents[1] / "scripts" / "excel-textconv.py"


def save_book(path, value="元の値", formula="=B1*2"):
    book = Workbook()
    sheet = book.active
    sheet.title = "売上"
    sheet["A1"] = value
    sheet["A2"] = formula
    sheet["B1"] = 10
    sheet = book.create_sheet("備考")
    sheet["C3"] = "日本語\n改行\tタブ"
    book.save(path)
    book.close()


def isolated_env(home):
    # Never let a developer's global Git configuration select another driver,
    # external diff command, pager, attributes file, or commit hook.
    env = {key: value for key, value in os.environ.items() if not key.startswith("GIT_")}
    env.update(
        HOME=str(home),
        XDG_CONFIG_HOME=str(home / ".config"),
        GIT_CONFIG_NOSYSTEM="1",
        GIT_CONFIG_GLOBAL=os.devnull,
        GIT_TERMINAL_PROMPT="0",
    )
    return env


class ConverterTestCase(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(prefix="excel-textconv-")
        self.addCleanup(directory.cleanup)
        self.directory = Path(directory.name)

    def convert(self, path, *, expected_returncode=0, env=None, options=()):
        result = subprocess.run(
            [sys.executable, str(CONVERTER), *options, str(path)],
            capture_output=True,
            encoding="utf-8",
            env=env,
            timeout=120,
        )
        self.assertEqual(result.returncode, expected_returncode, result.stderr)
        return result


class ConverterTests(ConverterTestCase):
    def test_japanese_sheets_formulas_and_multiline_cells(self):
        path = self.directory / "売上 表.xlsx"
        save_book(path)
        result = self.convert(path)
        self.assertEqual(
            result.stdout,
            'Sheet: "売上"\n'
            'A1\t"元の値"\n'
            'B1\t10\n'
            'A2\tformula: "=B1*2"\n'
            'Sheet: "備考"\n'
            'C3\t"日本語\\n改行\\tタブ"\n',
        )
        self.assertEqual(result.stderr, "")

    def test_extensionless_git_blob_and_ooxml_extensions(self):
        original = self.directory / "original.xlsx"
        save_book(original)
        expected = self.convert(original).stdout
        directory = self.directory / "日本語 と空白"
        directory.mkdir()
        for filename in ("git-blob", "book.xlsx", "book.xlsm", "book.xltx", "book.xltm"):
            with self.subTest(filename=filename):
                path = directory / filename
                shutil.copyfile(original, path)
                self.assertEqual(self.convert(path).stdout, expected)

    def test_dates_times_booleans_numbers_and_errors(self):
        path = self.directory / "types.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "型"
        for coordinate, value in {
            "A1": dt.datetime(2026, 10, 8, 15, 12, 3),
            "A2": dt.date(2026, 10, 8),
            "A3": dt.time(9, 30, 45),
            "A4": dt.timedelta(seconds=90),
            "A5": True,
            "A6": False,
            "A7": 42.0,
            "A8": 1.5,
            "A9": "#DIV/0!",
        }.items():
            sheet[coordinate] = value
        book.create_sheet("空シート")
        book.save(path)
        book.close()
        self.assertEqual(
            self.convert(path).stdout,
            'Sheet: "型"\n'
            'A1\tdatetime: 2026-10-08T15:12:03\n'
            'A2\tdatetime: 2026-10-08T00:00:00\n'
            'A3\ttime: 09:30:45\n'
            'A4\tduration-seconds: 90.0\n'
            'A5\ttrue\n'
            'A6\tfalse\n'
            'A7\t42\n'
            'A8\t1.5\n'
            'A9\terror: "#DIV/0!"\n'
            'Sheet: "空シート"\n',
        )

    def test_formatting_and_metadata_do_not_change_cell_output(self):
        original = self.directory / "original.xlsx"
        changed = self.directory / "formatted.xlsx"
        save_book(original)
        book = load_workbook(original)
        sheet = book["売上"]
        sheet["A1"].font = Font(bold=True, color="FF0000")
        sheet["A1"].fill = PatternFill("solid", fgColor="FFFF00")
        sheet.column_dimensions["A"].width = 50
        sheet.row_dimensions[1].height = 30
        sheet.freeze_panes = "B2"
        book.properties.creator = "別のユーザー"
        book.save(changed)
        book.close()
        self.assertNotEqual(original.read_bytes(), changed.read_bytes())
        self.assertEqual(self.convert(original).stdout, self.convert(changed).stdout)
        self.assertEqual(self.convert(changed).stdout, self.convert(changed).stdout)

    def test_incorrect_dimensions_do_not_hide_cells_or_expand_output(self):
        source = self.directory / "source.xlsx"
        book = Workbook()
        book.active["A1"] = "先頭"
        book.active["Z20"] = "末尾"
        book.save(source)
        book.close()
        original_bytes = source.read_bytes()
        expected = self.convert(source).stdout
        for dimension in ("A1:A1", "A1:XFD1048576"):
            with self.subTest(dimension=dimension):
                path = self.directory / "incorrect-dimension.xlsx"
                with zipfile.ZipFile(io.BytesIO(original_bytes)) as original:
                    with zipfile.ZipFile(path, "w") as changed:
                        for item in original.infolist():
                            contents = original.read(item.filename)
                            if item.filename == "xl/worksheets/sheet1.xml":
                                xml = ET.fromstring(contents)
                                node = xml.find(
                                    "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}dimension"
                                )
                                self.assertIsNotNone(node)
                                node.set("ref", dimension)
                                contents = ET.tostring(xml, encoding="utf-8")
                            changed.writestr(item, contents)
                self.assertEqual(self.convert(path).stdout, expected)

    def test_utf8_output_with_non_utf8_stdout_environment(self):
        path = self.directory / "日本語.xlsx"
        save_book(path)
        env = os.environ.copy()
        env["PYTHONIOENCODING"] = "ascii"
        self.assertIn('A1\t"元の値"', self.convert(path, env=env).stdout)

    def test_invalid_missing_and_corrupted_files_fail_clearly(self):
        invalid = self.directory / "not-a-workbook.xlsx"
        invalid.write_text("plain text", encoding="utf-8")
        corrupt = self.directory / "corrupted.xlsx"
        corrupt.write_bytes(b"PK\x03\x04truncated")
        for path in (invalid, corrupt, self.directory / "missing.xlsx"):
            with self.subTest(path=path.name):
                result = self.convert(path, expected_returncode=1)
                self.assertEqual(result.stdout, "")
                self.assertIn("Excel の差分変換に失敗しました:", result.stderr)
                self.assertNotIn("Traceback", result.stderr)


class CacheTests(ConverterTestCase):
    """--cache-dir: reuse results and finish slow conversions in the background."""

    def setUp(self):
        super().setUp()
        self.cache = self.directory / "キャッシュ 置き場"
        self.options = ("--cache-dir", str(self.cache))
        self.wait_env = os.environ.copy()
        self.wait_env["EXCEL_TEXTCONV_TIMEOUT"] = "0"
        self.slow_env = os.environ.copy()
        self.slow_env["EXCEL_TEXTCONV_TIMEOUT"] = "0.001"

    def key(self, path):
        return hashlib.sha256(path.read_bytes()).hexdigest()

    def hold_worker_lock(self):
        # A fresh lock means "a worker is running", so no real worker starts.
        self.cache.mkdir(parents=True, exist_ok=True)
        (self.cache / "worker.lock").touch()

    def run_worker(self):
        (self.cache / "worker.lock").unlink()
        subprocess.run(
            [sys.executable, str(CONVERTER), "--worker", *self.options], check=True, timeout=120
        )

    def wait_for(self, path):
        deadline = time.monotonic() + 120
        while not path.exists():
            self.assertLess(time.monotonic(), deadline, "background conversion did not finish")
            time.sleep(0.1)

    def test_cached_text_is_reused(self):
        path = self.directory / "売上.xlsx"
        save_book(path)
        expected = self.convert(path).stdout
        self.assertEqual(self.convert(path, env=self.wait_env, options=self.options).stdout, expected)
        cached = self.cache / (self.key(path) + ".txt")
        self.assertEqual(cached.read_text(encoding="utf-8"), expected)
        cached.write_text("cached\n", encoding="utf-8")
        self.assertEqual(self.convert(path, options=self.options).stdout, "cached\n")

    def test_slow_conversion_finishes_in_background_without_holding_git(self):
        path = self.directory / "大きい.xlsx"
        book = Workbook()
        for row in range(1, 2001):
            book.active.append([row * 10 + column for column in range(10)])
        book.save(path)
        book.close()
        expected = self.convert(path).stdout
        key = self.key(path)
        started = time.monotonic()
        result = self.convert(path, env=self.slow_env, options=self.options)
        self.assertLess(time.monotonic() - started, 15)
        self.assertEqual(
            result.stdout.splitlines()[0], "# Excel を変換中です (sha256 " + key[:16] + ")"
        )
        # The command returned while the detached worker was still converting.
        cached = self.cache / (key + ".txt")
        self.assertFalse(cached.exists())
        self.wait_for(cached)
        self.assertEqual(self.convert(path, options=self.options).stdout, expected)
        deadline = time.monotonic() + 30
        while (self.cache / "worker.lock").exists() and time.monotonic() < deadline:
            time.sleep(0.1)
        self.assertEqual(sorted(p.suffix for p in self.cache.iterdir() if p.suffix != ""), [".txt"])

    def test_queued_requests_answer_at_once_and_the_worker_converts_them(self):
        path = self.directory / "待ち.xlsx"
        save_book(path)
        self.hold_worker_lock()
        first = self.convert(path, env=self.slow_env, options=self.options).stdout
        self.assertTrue(first.startswith("# Excel を変換中です"))
        self.assertTrue(first.endswith('Sheet: "売上"\nSheet: "備考"\n'))
        self.assertTrue((self.cache / (self.key(path) + ".in")).exists())
        self.assertEqual(self.convert(path, env=self.slow_env, options=self.options).stdout, first)
        self.run_worker()
        self.assertFalse((self.cache / (self.key(path) + ".in")).exists())
        self.assertEqual(self.convert(path, options=self.options).stdout, self.convert(path).stdout)

    def test_background_failure_is_reported_like_a_direct_failure(self):
        path = self.directory / "破損.xlsx"
        path.write_bytes(b"PK\x03\x04truncated")
        self.hold_worker_lock()
        self.assertTrue(
            self.convert(path, env=self.slow_env, options=self.options).stdout.startswith(
                "# Excel を変換中です"
            )
        )
        self.run_worker()
        result = self.convert(path, expected_returncode=1, options=self.options)
        self.assertEqual(result.stdout, "")
        self.assertIn("Excel の差分変換に失敗しました:", result.stderr)
        self.assertNotIn("Traceback", result.stderr)

    def test_waits_for_the_result_when_the_timeout_is_zero(self):
        path = self.directory / "待つ.xlsx"
        save_book(path)
        self.hold_worker_lock()
        self.assertEqual(
            self.convert(path, env=self.wait_env, options=self.options).stdout,
            self.convert(path).stdout,
        )


@unittest.skipUnless(shutil.which("git"), "Git is required for textconv integration tests")
class GitTextconvTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(prefix="excel-git-diff-")
        self.addCleanup(directory.cleanup)
        self.directory = Path(directory.name)
        self.repo = self.directory / "repository"
        self.repo.mkdir()
        home = self.directory / "home"
        home.mkdir()
        self.env = isolated_env(home)
        converter_dir = self.directory / "変換 ツール"
        converter_dir.mkdir()
        converter = converter_dir / CONVERTER.name
        shutil.copyfile(CONVERTER, converter)
        self.git("init", "--quiet")
        self.git("config", "user.name", "Excel test")
        self.git("config", "user.email", "excel-test@example.invalid")
        self.git("config", "core.quotePath", "false")
        self.git("config", "core.autocrlf", "false")
        self.git("config", "diff.excel.textconv", shlex.join([sys.executable, str(converter)]))
        self.git("config", "diff.excel.binary", "true")
        (self.repo / ".gitattributes").write_text("*.xlsx diff=excel\n", encoding="utf-8")
        self.git("add", "--", ".gitattributes")
        self.git("commit", "--quiet", "-m", "Configure Excel textconv")

    def git(self, *args, expected_returncode=0):
        result = subprocess.run(
            ["git", "--no-pager", *args],
            cwd=self.repo,
            env=self.env,
            capture_output=True,
            encoding="utf-8",
            timeout=120,
        )
        self.assertEqual(result.returncode, expected_returncode, result.stderr)
        return result.stdout

    def commit_book(self, path):
        self.git("add", "--", path.name)
        self.git("commit", "--quiet", "-m", "Add workbook")

    def test_staged_and_unstaged_changes_compare_the_correct_blobs(self):
        path = self.repo / "売上 表.xlsx"
        save_book(path)
        self.commit_book(path)
        save_book(path, value="ステージ済み", formula="=B1*3")
        self.git("add", "--", path.name)
        save_book(path, value="作業中", formula="=B1*4")

        staged = self.git("diff", "--cached", "--textconv", "--", path.name)
        self.assertIn('-A1\t"元の値"', staged)
        self.assertIn('+A1\t"ステージ済み"', staged)
        self.assertIn('-A2\tformula: "=B1*2"', staged)
        self.assertIn('+A2\tformula: "=B1*3"', staged)
        self.assertNotIn("作業中", staged)

        unstaged = self.git("diff", "--textconv", "--", path.name)
        self.assertIn('-A1\t"ステージ済み"', unstaged)
        self.assertIn('+A1\t"作業中"', unstaged)
        self.assertIn('-A2\tformula: "=B1*3"', unstaged)
        self.assertIn('+A2\tformula: "=B1*4"', unstaged)
        self.assertNotIn("元の値", unstaged)
        self.assertNotIn("Binary files", staged + unstaged)

    def test_untracked_staged_added_and_deleted_workbooks(self):
        path = self.repo / "新規 ブック.xlsx"
        save_book(path)
        untracked = self.git(
            "diff", "--no-index", "--textconv", "--", os.devnull, path.name,
            expected_returncode=1,
        )
        self.assertIn('+Sheet: "売上"', untracked)
        self.assertIn('+A1\t"元の値"', untracked)
        self.git("add", "--", path.name)
        added = self.git("diff", "--cached", "--textconv", "--", path.name)
        self.assertIn('+Sheet: "備考"', added)
        self.assertIn('+C3\t"日本語\\n改行\\tタブ"', added)
        self.git("commit", "--quiet", "-m", "Add workbook")

        path.unlink()
        deleted = self.git("diff", "--textconv", "--", path.name)
        self.assertIn('-Sheet: "売上"', deleted)
        self.assertIn('-A1\t"元の値"', deleted)
        self.git("add", "-u", "--", path.name)
        staged_deleted = self.git("diff", "--cached", "--textconv", "--", path.name)
        self.assertIn('-A2\tformula: "=B1*2"', staged_deleted)
        self.assertNotIn("Binary files", untracked + added + deleted + staged_deleted)

    def test_cached_textconv_gives_the_same_diff(self):
        path = self.repo / "キャッシュ.xlsx"
        save_book(path)
        self.commit_book(path)
        save_book(path, value="変更後")
        expected = self.git("diff", "--textconv", "--", path.name)
        self.env["EXCEL_TEXTCONV_TIMEOUT"] = "0"
        cache = self.directory / "変換 キャッシュ"
        command = self.git("config", "diff.excel.textconv").strip()
        self.git("config", "diff.excel.textconv", command + " --cache-dir " + shlex.quote(str(cache)))
        for _ in range(2):
            self.assertEqual(self.git("diff", "--textconv", "--", path.name), expected)
        self.assertEqual(len(list(cache.glob("*.txt"))), 2)

        # Only the old side is converted; the new one is still queued.
        save_book(path, value="未変換")
        (cache / "worker.lock").touch()
        self.env["EXCEL_TEXTCONV_TIMEOUT"] = "0.001"
        changed = [
            line for line in self.git("diff", "--textconv", "--", path.name).splitlines()
            if line[:1] in "+-" and not line.startswith(("+++", "---"))
        ]
        self.assertTrue(changed[0].startswith("+# Excel を変換中です"), changed[:3])
        self.assertIn('-A1\t"元の値"', changed)

    def test_formatting_only_change_has_no_cell_diff(self):
        path = self.repo / "formatting.xlsx"
        save_book(path)
        self.commit_book(path)
        original = path.read_bytes()
        book = load_workbook(path)
        book["売上"]["A1"].font = Font(bold=True)
        book.save(path)
        book.close()
        self.assertNotEqual(path.read_bytes(), original)
        self.assertEqual(self.git("diff", "--textconv", "--", path.name), "")
        self.git("add", "--", path.name)
        self.assertEqual(self.git("diff", "--cached", "--textconv", "--", path.name), "")


if __name__ == "__main__":
    unittest.main()
