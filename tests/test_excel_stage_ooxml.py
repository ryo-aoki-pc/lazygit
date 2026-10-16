"""Verify partial Excel edits preserve values and unsupported package content."""

import datetime as dt
import importlib.util
import io
import json
import random
from decimal import localcontext
from pathlib import Path
import sys
import unittest
from unittest import mock
import zipfile

from lxml import etree
from openpyxl import Workbook as OpenpyxlWorkbook, load_workbook
from openpyxl.styles import Font


MODULE_PATH = Path(__file__).resolve().parents[1] / "scripts/excel_stage_ooxml.py"
SPEC = importlib.util.spec_from_file_location("excel_stage_ooxml_unit", MODULE_PATH)
OOXML = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = OOXML
SPEC.loader.exec_module(OOXML)
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def book_bytes():
    book = OpenpyxlWorkbook()
    sheet = book.active
    sheet.title = "売上"
    sheet["A1"] = "インデックス"
    sheet["A1"].font = Font(bold=True)
    sheet["B1"] = 100
    sheet["B2"] = "=B1*2"
    sheet["C1"] = "残す"
    sheet["E1"] = dt.datetime(2026, 10, 8)
    sheet.row_dimensions[1].height = 31
    book.create_sheet("備考")["A1"] = "別シート"
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def package(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def replace_parts(data, replacements):
    parts = package(data)
    parts.update(replacements)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return output.getvalue()


def edit_xml(data, path, edit):
    parts = package(data)
    root = etree.fromstring(parts[path])
    edit(root)
    return replace_parts(data, {path: etree.tostring(root)})


def edit_book(data, edit):
    book = load_workbook(io.BytesIO(data))
    edit(book)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def cell_node(root, coordinate):
    return root.find(".//{" + NS + "}c[@r='" + coordinate + "']")


class OoxmlStagingTests(unittest.TestCase):
    def test_number_fast_path_matches_decimal_at_normalization_boundaries(self):
        cases = {
            "0": "0", "-0": "0", "+0.000e30": "0", "0000": "0",
            "1.230000": "1.23", "-000001.230000": "-1.23",
            "0.000001000": "0.000001", "0.0000001000": "1E-7",
            "-0.000001000": "-0.000001", "-0.0000001000": "-1E-7",
            "999999999999999": "999999999999999",
            "9999999999999999": "9999999999999999",
            "1000000000000000": "1000000000000000",
            "10000000000000000": "1E+16",
            "-10000000000000000": "-1E+16",
            "12345678901234567": "12345678901234567",
            "1234567890123456789.1": "1234567890123456789.1",
            "10000000000000001": "10000000000000001",
            "10000000000000010": "1.000000000000001E+16",
            "1.23456789012345678901234567890": "1.2345678901234567890123456789",
            "1234500e-10": "0.00012345", "1234500E+10": "1.2345E+16",
            ".1": "0.1", "1.": "1", "+000001e+000015": "1000000000000000",
            "1e999999": "1E+999999", "1e-999999": "1E-999999",
        }
        with localcontext() as context:
            context.prec = 6  # Normalization must retain digits beyond context precision.
            for value, expected in cases.items():
                with self.subTest(value=value):
                    self.assertEqual(OOXML._number(value), expected)
                    self.assertEqual(OOXML._number(value), OOXML._decimal_number(value))

    def test_number_fast_path_matches_decimal_for_random_numeric_spellings(self):
        randomizer = random.Random(692750)
        for index in range(5000):
            digits = "".join(str(randomizer.randrange(10)) for _ in range(randomizer.randrange(1, 81)))
            point = randomizer.randrange(len(digits) + 1)
            value = randomizer.choice(("", "+", "-")) + "0" * randomizer.randrange(5)
            value += digits if index % 3 == 0 else digits[:point] + "." + digits[point:]
            if index % 2:
                value += randomizer.choice(("e", "E")) + randomizer.choice(("", "+")) + str(randomizer.randrange(101))
            elif index % 3:
                value += "e-" + str(randomizer.randrange(101))
            with self.subTest(value=value):
                self.assertEqual(OOXML._number(value), OOXML._decimal_number(value))

    def test_number_decimal_fallback_keeps_additional_spellings_and_errors(self):
        for value in (" 1.2300 ", "1_000.00", "\u0661\u0662.\u0663\u0660", "1e0000000000000016", "1e1000000", 0.1, 12):
            with self.subTest(value=value):
                self.assertEqual(OOXML._number(value), OOXML._decimal_number(value))
        for value in ("", ".", "++1", "1e", "1..0", "NaN", "sNaN", "Infinity", "-inf", "0e999999999999999999999"):
            with self.subTest(value=value), self.assertRaises(OOXML.WorkbookError):
                OOXML._number(value)

    def test_cell_record_fast_path_matches_json_dumps_exactly(self):
        class UnhashableString(str):
            __hash__ = None

        values = (
            "", "日本語", 'quote " backslash \\', "/", "\b\f\n\r\t",
            "".join(chr(value) for value in range(32)), "\x7f\x85\u2028\u2029",
            "\U0001f600", "\ud800", "\udfff", "\ud800\udfff",
            True, False, None, 0, -2, 0.125, ["日本語", None], {"引用符\"": "\\"},
        )
        for coordinate in ("A1", "AA2", "XFD1048576"):
            for kind in ("string", "number", "boolean", "formula", "date", "error", "特殊\"\\\n", UnhashableString("string")):
                for value in values:
                    with self.subTest(coordinate=coordinate, kind=kind, value=repr(value)):
                        expected = json.dumps([coordinate, kind, value], ensure_ascii=False, allow_nan=False, separators=(",", ":"))
                        self.assertEqual(OOXML.cell_record(coordinate, OOXML.Cell(kind, value)), expected)
            self.assertEqual(OOXML.cell_record(coordinate, None), json.dumps([coordinate, "blank", None], ensure_ascii=False, allow_nan=False, separators=(",", ":")))
        for value in (float("nan"), float("inf"), float("-inf")):
            with self.subTest(value=value), self.assertRaises(ValueError):
                OOXML.cell_record("A1", OOXML.Cell("number", value))
        with self.assertRaises(OOXML.WorkbookError):
            OOXML.cell_record("A0", OOXML.Cell("string", "invalid coordinate"))

    def test_projection_orders_cells_by_numeric_row_and_column(self):
        base = OOXML.Workbook(book_bytes())

        def change(book):
            for coordinate in ("A10", "AA2", "B2", "A2"):
                book["売上"][coordinate] = coordinate

        source = OOXML.Workbook(edit_book(base.data, change))
        records = [OOXML.parse_record(line)[0] for line in source.projection(base)["1"].splitlines()]
        self.assertEqual([coordinate for coordinate in records if coordinate in {"A2", "B2", "AA2", "A10"}], ["A2", "B2", "AA2", "A10"])

    def test_streaming_load_rejects_invalid_cells_duplicate_elements_and_dtd(self):
        original = book_bytes()
        invalid = (
            '<sheetData><row r="1"><c r="A0"><v>1</v></c></row></sheetData>',
            '<sheetData><row r="1"><c r="XFE1"><v>1</v></c></row></sheetData>',
            '<sheetData><row r="1"><c r="A1"/><c r="A1"><v>1</v></c></row></sheetData>',
            '<sheetData><row r="1"><c r="A1"><v>1</v><v>2</v></c></row></sheetData>',
            '<sheetData><row r="1"><c r="A1"><f>1</f><f>2</f></c></row></sheetData>',
            '<sheetData><row r="1"><c r="A1" t="inlineStr"><is/><is/></c></row></sheetData>',
            '<sheetData><row r="1"><c r="A1" t="s"><v>-1</v></c></row></sheetData>',
            '<sheetData><row r="1"><c r="A1" t="s"><v>0</v></c></row></sheetData>',
            '<sheetData><row r="1"><c r="A1" t="b"><v>2</v></c></row></sheetData>',
            '<sheetData/><sheetData/>',
            '<sheetData/><mergeCells/><mergeCells/>',
        )
        for content in invalid:
            with self.subTest(content=content), self.assertRaises(OOXML.WorkbookError):
                worksheet = ('<worksheet xmlns="' + NS + '">' + content + '</worksheet>').encode("utf-8")
                OOXML.Workbook(replace_parts(original, {"xl/worksheets/sheet1.xml": worksheet}))
        worksheet = ('<!DOCTYPE worksheet [<!ENTITY value "1">]>'
                     '<worksheet xmlns="' + NS + '"><sheetData><row r="1">'
                     '<c r="A1"><v>&value;</v></c></row></sheetData></worksheet>').encode("utf-8")
        with self.assertRaises(OOXML.WorkbookError):
            OOXML.Workbook(replace_parts(original, {"xl/worksheets/sheet1.xml": worksheet}))

    def test_streaming_load_keeps_unsafe_cells_and_trailing_merge_checks(self):
        original = book_bytes()
        cases = (
            '<c r="A1" cm="1"><v>1</v></c>',
            '<c r="A1"><v>1</v><extLst/></c>',
            '<c r="A1" t="unknown"><v>1</v></c>',
            '<c r="A1"><v>1</v><x:v xmlns:x="urn:extra">2</x:v></c>',
            '<c r="A1"><f t="shared" ref="A1:B2">1</f><v>1</v></c>',
            '<c r="A1" t="inlineStr"><is><r><t>rich</t></r></is></c>',
        )
        for content in cases:
            with self.subTest(content=content):
                worksheet = ('<worksheet xmlns="' + NS + '"><sheetData><row r="1">' + content + '</row></sheetData></worksheet>').encode("utf-8")
                base = OOXML.Workbook(replace_parts(original, {"xl/worksheets/sheet1.xml": worksheet}))
                with self.assertRaises(OOXML.WorkbookError):
                    base.rebuild(base, {("1", "A1")})
        worksheet = ('<worksheet xmlns="' + NS + '"><sheetData><row r="1">'
                     '<c r="A1"><v>1</v></c><c r="B1"><v>2</v></c></row></sheetData>'
                     '<mergeCells><mergeCell ref="A1:B1"/></mergeCells></worksheet>').encode("utf-8")
        base = OOXML.Workbook(replace_parts(original, {"xl/worksheets/sheet1.xml": worksheet}))
        with self.assertRaises(OOXML.WorkbookError):
            base.rebuild(base, {("1", "B1")})

    def test_streaming_load_sorts_unsorted_xml_and_keeps_unusual_namespaces(self):
        worksheet = ('<worksheet xmlns="' + NS + '" xmlns:x="urn:row" xmlns:y="urn:cell">'
                     '<sheetData><row r="10"><c r="A10"><v>10</v></c></row>'
                     '<x:row r="2"><c r="AA2"><v>3</v></c><y:c r="B2"><y:v>2</y:v></y:c>'
                     '<c r="A2"><v>1</v></c></x:row></sheetData></worksheet>').encode("utf-8")
        base = OOXML.Workbook(replace_parts(book_bytes(), {"xl/worksheets/sheet1.xml": worksheet}))
        records = [OOXML.parse_record(line) for line in base.projection()["1"].splitlines()]
        self.assertEqual(records, [(coordinate, OOXML.Cell("number", value)) for coordinate, value in (
            ("A2", "1"), ("B2", "2"), ("AA2", "3"), ("A10", "10"),
        )])
        with self.assertRaises(OOXML.WorkbookError):
            base.rebuild(base, {("1", "B2")})
        self.assertTrue(all(sheet.path not in base._trees for sheet in base.sheets))

    def test_projection_tracks_public_cells_dictionary_key_and_value_changes(self):
        base = OOXML.Workbook(book_bytes())
        count = len(base.cells)
        base.cells[("1", "AA2")] = base.cells.pop(("1", "A1"))
        base.cells[("1", "B1")] = OOXML.Cell("number", "3")
        self.assertEqual(len(base.cells), count)
        records = [OOXML.parse_record(line) for line in base.projection()["1"].splitlines()]
        expected = [(coordinate, cell) for (identifier, coordinate), cell in sorted(
            base.cells.items(), key=lambda item: (item[0][0], OOXML.coordinate_key(item[0][1]))
        ) if identifier == "1"]
        self.assertEqual(records, expected)
        base.cells[("1", "A10")] = OOXML.Cell("boolean", True)
        self.assertIn('["A10","boolean",true]\n', base.projection()["1"])
        del base.cells[("1", "AA2")]
        self.assertNotIn('["AA2",', base.projection()["1"])

    def test_rebuild_verification_rejects_wrong_missing_added_and_unselected_cells(self):
        base = OOXML.Workbook(book_bytes())
        work = OOXML.Workbook(edit_book(base.data, lambda book: setattr(book["売上"]["A1"], "value", "changed")))
        original_replace = OOXML.Workbook._replace
        for corruption in ("wrong", "missing", "added", "unselected"):
            def broken_replace(book, data, rows, coordinate, cell):
                if corruption == "wrong":
                    return
                original_replace(book, data, rows, coordinate, None if corruption == "missing" else cell)
                if corruption == "added":
                    original_replace(book, data, rows, "Z99", OOXML.Cell("number", "1"))
                elif corruption == "unselected":
                    original_replace(book, data, rows, "C1", OOXML.Cell("string", "corrupt"))

            with self.subTest(corruption=corruption), mock.patch.object(OOXML.Workbook, "_replace", broken_replace):
                with self.assertRaises(OOXML.WorkbookError):
                    base.rebuild(work, {("1", "A1")})

    def test_rebuild_verifies_changed_parts_only_and_keeps_other_sheet_bytes(self):
        base = OOXML.Workbook(book_bytes())
        work = OOXML.Workbook(edit_book(base.data, lambda book: setattr(book["売上"]["A1"], "value", "changed")))
        checked = []
        original_load = OOXML.Workbook._load_sheet

        def load_sheet(book, sheet, strings, expected=None):
            checked.append(sheet.id)
            return original_load(book, sheet, strings, expected)

        with mock.patch.object(OOXML.Workbook, "_load_sheet", load_sheet):
            result = base.rebuild(work, {("1", "A1")})
        self.assertEqual(checked, ["1"])
        self.assertEqual(package(base.data)["xl/worksheets/sheet2.xml"], package(result)["xl/worksheets/sheet2.xml"])

    def test_rebuild_verification_rejects_changes_to_an_unselected_sheet_alias(self):
        def alias(root):
            sheets = list(root)
            sheets[1].set("Target", sheets[0].get("Target"))

        original = edit_xml(book_bytes(), "xl/_rels/workbook.xml.rels", alias)
        base = OOXML.Workbook(original)
        work_data = edit_xml(original, "xl/worksheets/sheet1.xml", lambda root: setattr(cell_node(root, "B1").find("{" + NS + "}v"), "text", "200"))
        work = OOXML.Workbook(work_data)
        with self.assertRaises(OOXML.WorkbookError):
            base.rebuild(work, {("1", "B1")})
        result = OOXML.Workbook(base.rebuild(work, {("1", "B1"), ("2", "B1")}))
        self.assertEqual(result.cells[("1", "B1")], OOXML.Cell("number", "200"))
        self.assertEqual(result.cells[("2", "B1")], OOXML.Cell("number", "200"))

    def test_rebuild_verification_keeps_structure_and_checks_untouched_zip_bytes(self):
        base = OOXML.Workbook(book_bytes())
        work = OOXML.Workbook(edit_book(base.data, lambda book: setattr(book["売上"]["A1"], "value", "changed")))
        original_recalculate = OOXML.Workbook._recalculate

        def change_structure(book, updates, removed):
            original_recalculate(book, updates, removed)
            root = etree.fromstring(updates[book._workbook_path])
            root.find("{" + NS + "}sheets/{" + NS + "}sheet").set("name", "corrupt")
            updates[book._workbook_path] = etree.tostring(root)

        with mock.patch.object(OOXML.Workbook, "_recalculate", change_structure):
            with self.assertRaises(OOXML.WorkbookError):
                base.rebuild(work, {("1", "A1")})
        for corruption in ("remove", "change"):
            def corrupt_other_sheet(book, updates, removed):
                original_recalculate(book, updates, removed)
                other_path = book.sheets[1].path
                if corruption == "remove":
                    removed.add(other_path)
                else:
                    root = etree.fromstring(book._parts[other_path])
                    root.find(".//{" + NS + "}t").text = "corrupt"
                    updates[other_path] = etree.tostring(root)

            with self.subTest(corruption=corruption), mock.patch.object(OOXML.Workbook, "_recalculate", corrupt_other_sheet):
                with self.assertRaises(OOXML.WorkbookError):
                    base.rebuild(work, {("1", "A1")})
        original_write = zipfile.ZipFile.writestr

        def corrupt_part(archive, info, content, *args, **kwargs):
            if info.filename == "xl/worksheets/sheet2.xml":
                content += b"<!-- corrupt -->"
            return original_write(archive, info, content, *args, **kwargs)

        with mock.patch.object(zipfile.ZipFile, "writestr", corrupt_part):
            with self.assertRaises(OOXML.WorkbookError):
                base.rebuild(work, {("1", "A1")})

    def test_selected_cells_preserve_index_styles_other_cells_and_package_parts(self):
        original = replace_parts(book_bytes(), {
            "xl/vbaProject.bin": b"opaque VBA bytes\x00\xff",
            "customXml/item1.xml": b"<unknown xmlns='urn:private'>preserve me</unknown>",
        })

        def change(book):
            book["売上"]["A1"] = "作業中\n改行\tタブ"
            book["売上"]["A1"].font = Font(italic=True)
            book["売上"]["C1"] = "未選択"
            book["売上"].row_dimensions[1].height = 45
            book["備考"]["A1"] = "未選択の別シート"

        source = edit_book(original, change)
        base = OOXML.Workbook(original)
        result = base.rebuild(OOXML.Workbook(source), {("1", "A1")})
        book = load_workbook(io.BytesIO(result))
        self.assertEqual(book["売上"]["A1"].value, "作業中\n改行\tタブ")
        self.assertTrue(book["売上"]["A1"].font.bold)
        self.assertFalse(book["売上"]["A1"].font.italic)
        self.assertEqual(book["売上"].row_dimensions[1].height, 31)
        self.assertEqual(book["売上"]["C1"].value, "残す")
        self.assertEqual(book["備考"]["A1"].value, "別シート")
        book.close()
        before, after = package(original), package(result)
        modified = {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}
        self.assertEqual(set(before), set(after))
        for name in set(before) - modified:
            self.assertEqual(before[name], after[name], name)

    def test_projection_uses_explicit_blanks_and_preserves_empty_string_types(self):
        base = OOXML.Workbook(book_bytes())

        def change(book):
            book["売上"]["A1"] = None
            book["売上"]["D4"] = "空文字にする"
            book["売上"]["F1"] = False
            book["売上"]["G1"] = 0
            book["売上"]["H1"] = "#DIV/0!"

        source = edit_book(base.data, change)

        def empty(root):
            node = cell_node(root, "D4")
            node.find("{" + NS + "}is/{" + NS + "}t").text = ""

        source = OOXML.Workbook(edit_xml(source, "xl/worksheets/sheet1.xml", empty))
        records = [OOXML.parse_record(line) for line in source.projection(base)["1"].splitlines()]
        self.assertIn(("A1", None), records)
        self.assertIn(("D4", OOXML.Cell("string", "")), records)
        self.assertIn(("F1", OOXML.Cell("boolean", False)), records)
        self.assertIn(("G1", OOXML.Cell("number", "0")), records)
        self.assertIn(("H1", OOXML.Cell("error", "#DIV/0!")), records)
        self.assertIn('["D4","blank",null]', base.projection(source)["1"])
        selected = {("1", value) for value in ("A1", "D4", "F1", "G1", "H1")}
        actual = OOXML.Workbook(base.rebuild(source, selected))
        self.assertNotIn(("1", "A1"), actual.cells)
        for key in selected - {("1", "A1")}:
            self.assertEqual(actual.cells[key], source.cells[key])

    def test_shared_strings_are_resolved_without_copying_source_string_indices(self):
        base = OOXML.Workbook(book_bytes())
        source = base.data

        def shared_cell(root):
            node = cell_node(root, "A1")
            for child in list(node):
                node.remove(child)
            node.set("t", "s")
            etree.SubElement(node, "{" + NS + "}v").text = "1"

        source = edit_xml(source, "xl/worksheets/sheet1.xml", shared_cell)

        def relationship(root):
            etree.SubElement(root, "{" + REL_NS + "}Relationship", Id="rIdStrings", Type=DOC_REL_NS + "/sharedStrings", Target="sharedStrings.xml")

        source = edit_xml(source, "xl/_rels/workbook.xml.rels", relationship)
        source = replace_parts(source, {"xl/sharedStrings.xml": (
            '<sst xmlns="' + NS + '"><si><t>別の文字列</t></si>'
            '<si><t xml:space="preserve"> 日本語\n共有文字列 </t></si></sst>'
        ).encode("utf-8")})
        work = OOXML.Workbook(source)
        result = base.rebuild(work, {("1", "A1")})
        self.assertEqual(OOXML.Workbook(result).cells[("1", "A1")], OOXML.Cell("string", " 日本語\n共有文字列 "))
        self.assertNotIn("xl/sharedStrings.xml", package(result))
        node = cell_node(etree.fromstring(package(result)["xl/worksheets/sheet1.xml"]), "A1")
        self.assertEqual(node.get("t"), "inlineStr")

    def test_phonetic_readings_are_plain_text_but_formatted_runs_are_rejected(self):
        # Japanese Excel keeps readings (rPh) and their settings (phoneticPr)
        # beside the text of ordinary shared and inline strings.
        strings = (
            '<sst xmlns="' + NS + '">'
            '<si><t>売上</t><rPh sb="0" eb="2"><t>ウリアゲ</t></rPh><phoneticPr fontId="1"/></si>'
            '<si><t>大阪</t><rPh sb="0" eb="2"><t>オオサカ</t></rPh><phoneticPr fontId="1"/></si>'
            '<si><r><rPr><b/></rPr><t>太字</t></r></si></sst>'
        ).encode("utf-8")

        def book(index, inline_text):
            def cells(root):
                node = cell_node(root, "A1")
                for child in list(node):
                    node.remove(child)
                node.set("t", "s")
                etree.SubElement(node, "{" + NS + "}v").text = str(index)
                inline = cell_node(root, "C1").find("{" + NS + "}is")
                inline.find("{" + NS + "}t").text = inline_text
                etree.SubElement(inline, "{" + NS + "}phoneticPr", fontId="1")

            def relationship(root):
                etree.SubElement(root, "{" + REL_NS + "}Relationship", Id="rIdStrings", Type=DOC_REL_NS + "/sharedStrings", Target="sharedStrings.xml")

            data = edit_xml(book_bytes(), "xl/worksheets/sheet1.xml", cells)
            data = edit_xml(data, "xl/_rels/workbook.xml.rels", relationship)
            return OOXML.Workbook(replace_parts(data, {"xl/sharedStrings.xml": strings}))

        base = book(0, "東京")
        self.assertEqual(base.cells[("1", "A1")], OOXML.Cell("string", "売上"))
        result = OOXML.Workbook(base.rebuild(book(1, "名古屋"), {("1", "A1"), ("1", "C1")}))
        self.assertEqual(result.cells[("1", "A1")], OOXML.Cell("string", "大阪"))
        self.assertEqual(result.cells[("1", "C1")], OOXML.Cell("string", "名古屋"))
        with self.assertRaises(OOXML.WorkbookError):
            base.rebuild(book(2, "東京"), {("1", "A1")})

    def test_number_format_and_formula_cache_changes_do_not_change_cell_values(self):
        original = book_bytes()

        def cache(root):
            cell_node(root, "B2").find("{" + NS + "}v").text = "200"
            cell_node(root, "B1").find("{" + NS + "}v").text = "100.00"

        original = edit_xml(original, "xl/worksheets/sheet1.xml", cache)
        base = OOXML.Workbook(original)

        def change(root):
            cell_node(root, "B2").find("{" + NS + "}v").text = "999"
            cell_node(root, "E1").set("s", "0")

        source = OOXML.Workbook(edit_xml(original, "xl/worksheets/sheet1.xml", change))
        self.assertEqual(base.cells, source.cells)
        self.assertEqual(base.projection(source), source.projection(base))
        self.assertEqual(base.rebuild(source, set()), original)
        self.assertEqual(base.cells[("1", "B1")], OOXML.Cell("number", "100"))
        self.assertEqual(base.cells[("1", "B2")], OOXML.Cell("formula", "=B1*2"))

    def test_formula_edit_clears_cache_removes_calculation_chain_and_keeps_manual_mode(self):
        original = book_bytes()

        def cache(root):
            cell_node(root, "B2").find("{" + NS + "}v").text = "200"

        original = edit_xml(original, "xl/worksheets/sheet1.xml", cache)

        def calculation(root):
            root.find("{" + NS + "}calcPr").set("calcMode", "manual")

        original = edit_xml(original, "xl/workbook.xml", calculation)

        def relationship(root):
            etree.SubElement(root, "{" + REL_NS + "}Relationship", Id="rIdChain", Type=DOC_REL_NS + "/calcChain", Target="calcChain.xml")

        original = edit_xml(original, "xl/_rels/workbook.xml.rels", relationship)

        def override(root):
            etree.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", PartName="/xl/calcChain.xml", ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml")

        original = edit_xml(original, "[Content_Types].xml", override)
        original = replace_parts(original, {"xl/calcChain.xml": ('<calcChain xmlns="' + NS + '"><c r="B2" i="1"/></calcChain>').encode()})
        source = edit_xml(original, "xl/worksheets/sheet1.xml", lambda root: setattr(cell_node(root, "B2").find("{" + NS + "}f"), "text", "B1*3"))
        result = OOXML.Workbook(original).rebuild(OOXML.Workbook(source), {("1", "B2")})
        parts = package(result)
        self.assertNotIn("xl/calcChain.xml", parts)
        self.assertNotIn(b"calcChain", parts["xl/_rels/workbook.xml.rels"])
        self.assertNotIn(b"calcChain", parts["[Content_Types].xml"])
        formula = cell_node(etree.fromstring(parts["xl/worksheets/sheet1.xml"]), "B2")
        self.assertIsNone(formula.find("{" + NS + "}v"))
        self.assertEqual(formula.find("{" + NS + "}f").text, "B1*3")
        calculation = etree.fromstring(parts["xl/workbook.xml"]).find("{" + NS + "}calcPr")
        self.assertEqual(calculation.get("calcMode"), "manual")
        self.assertEqual(calculation.get("fullCalcOnLoad"), "1")
        # A persistent "always fully recalculate" setting must not be added.
        self.assertIsNone(calculation.get("forceFullCalc"))

    def test_namespace_declarations_and_untouched_rich_cells_survive(self):
        parts = package(book_bytes())
        worksheet = parts["xl/worksheets/sheet1.xml"].replace(
            ('xmlns="' + NS + '"').encode(),
            ('xmlns="' + NS + '" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac" mc:Ignorable="x14ac"').encode(), 1,
        )
        original = replace_parts(book_bytes(), {"xl/worksheets/sheet1.xml": worksheet})

        def extension(root):
            node = cell_node(root, "C1")
            node.set("cm", "1")
            etree.SubElement(node, "{" + NS + "}extLst")

        original = edit_xml(original, "xl/worksheets/sheet1.xml", extension)
        source = edit_xml(original, "xl/worksheets/sheet1.xml", lambda root: setattr(cell_node(root, "A1").find("{" + NS + "}is/{" + NS + "}t"), "text", "選択値"))
        result = OOXML.Workbook(original).rebuild(OOXML.Workbook(source), {("1", "A1")})
        before = etree.fromstring(package(original)["xl/worksheets/sheet1.xml"])
        after = etree.fromstring(package(result)["xl/worksheets/sheet1.xml"])
        self.assertEqual(before.nsmap, after.nsmap)
        self.assertEqual(before.get("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable"), after.get("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable"))
        self.assertEqual(etree.tostring(cell_node(before, "C1")), etree.tostring(cell_node(after, "C1")))
        with self.assertRaises(OOXML.WorkbookError):
            OOXML.Workbook(original).rebuild(OOXML.Workbook(source), {("1", "C1")})

    def test_unsupported_formula_groups_and_merged_nonanchor_changes_are_rejected(self):
        for formula_type in ("shared", "array", "dataTable"):
            with self.subTest(formula_type=formula_type):
                original = edit_xml(book_bytes(), "xl/worksheets/sheet1.xml", lambda root: cell_node(root, "B2").find("{" + NS + "}f").set("t", formula_type))
                work = OOXML.Workbook(original)
                with self.assertRaises(OOXML.WorkbookError):
                    work.rebuild(work, {("1", "B2")})

        def merge(root):
            group = etree.SubElement(root, "{" + NS + "}mergeCells")
            etree.SubElement(group, "{" + NS + "}mergeCell", ref="B1:C1")

        base = OOXML.Workbook(edit_xml(book_bytes(), "xl/worksheets/sheet1.xml", merge))
        with self.assertRaisesRegex(OOXML.WorkbookError, "結合セル"):
            base.rebuild(base, {("1", "C1")})

    def test_structure_changes_and_malformed_selection_are_rejected(self):
        base = OOXML.Workbook(book_bytes())
        renamed = edit_book(base.data, lambda book: setattr(book["売上"], "title", "別名"))
        with self.assertRaises(OOXML.WorkbookError):
            base.assert_compatible(OOXML.Workbook(renamed))
        epoch = edit_xml(base.data, "xl/workbook.xml", lambda root: root.find("{" + NS + "}workbookPr").set("date1904", "1"))
        with self.assertRaises(OOXML.WorkbookError):
            base.assert_compatible(OOXML.Workbook(epoch))
        malformed = ('not JSON', '["A1",[],null]', '["A1","unknown",null]', '["A0","number","1"]', '["A1","boolean",1]', '["A1","number","NaN"]')
        for line in malformed:
            with self.subTest(line=line), self.assertRaises(OOXML.WorkbookError):
                OOXML.parse_record(line)
        with self.assertRaises(OOXML.WorkbookError):
            OOXML.Workbook(b"PK truncated")

    def test_external_workbook_target_changes_are_rejected_but_cached_values_are_ignored(self):
        original = book_bytes()

        def relationship(root):
            etree.SubElement(root, "{" + REL_NS + "}Relationship", Id="rIdExternal", Type=DOC_REL_NS + "/externalLink", Target="externalLinks/externalLink1.xml")

        original = edit_xml(original, "xl/_rels/workbook.xml.rels", relationship)
        original = replace_parts(original, {
            "xl/externalLinks/externalLink1.xml": ('<externalLink xmlns="' + NS + '" xmlns:r="' + DOC_REL_NS + '"><externalBook r:id="rId1"><sheetNames><sheetName val="Data"/></sheetNames><sheetDataSet/></externalBook></externalLink>').encode(),
            "xl/externalLinks/_rels/externalLink1.xml.rels": ('<Relationships xmlns="' + REL_NS + '"><Relationship Id="rId1" Type="' + DOC_REL_NS + '/externalLinkPath" Target="old.xlsx" TargetMode="External"/></Relationships>').encode(),
        })
        base = OOXML.Workbook(original)
        changed_target = edit_xml(original, "xl/externalLinks/_rels/externalLink1.xml.rels", lambda root: root[0].set("Target", "different.xlsx"))
        with self.assertRaises(OOXML.WorkbookError):
            base.assert_compatible(OOXML.Workbook(changed_target))

        def cached_value(root):
            sheet = etree.SubElement(root[0].find("{" + NS + "}sheetDataSet"), "{" + NS + "}sheetData", sheetId="0")
            row = etree.SubElement(sheet, "{" + NS + "}row", r="1")
            cell = etree.SubElement(row, "{" + NS + "}cell", r="A1", t="n")
            etree.SubElement(cell, "{" + NS + "}v").text = "999"

        changed_cache = edit_xml(original, "xl/externalLinks/externalLink1.xml", cached_value)
        base.assert_compatible(OOXML.Workbook(changed_cache))

    def test_inserted_rows_expand_dimension_and_date_type_is_retained(self):
        base = OOXML.Workbook(book_bytes())
        source = edit_book(base.data, lambda book: setattr(book["売上"]["Z99"], "value", dt.datetime(2027, 1, 2)))

        def iso_date(root):
            node = cell_node(root, "Z99")
            node.set("t", "d")
            node.find("{" + NS + "}v").text = "2027-01-02T00:00:00"

        source = OOXML.Workbook(edit_xml(source, "xl/worksheets/sheet1.xml", iso_date))
        result = base.rebuild(source, {("1", "Z99")})
        actual = OOXML.Workbook(result)
        self.assertEqual(actual.cells[("1", "Z99")], OOXML.Cell("date", "2027-01-02T00:00:00"))
        root = etree.fromstring(package(result)["xl/worksheets/sheet1.xml"])
        self.assertEqual(root.find("{" + NS + "}dimension").get("ref"), "A1:Z99")
        book = load_workbook(io.BytesIO(result))
        self.assertEqual(book["売上"]["Z99"].value, dt.datetime(2027, 1, 2))
        book.close()

    def test_many_selected_cells_across_existing_and_new_rows_keep_xml_order(self):
        def change(book):
            sheet = book["売上"]
            for row in range(1, 301):
                sheet.cell(row, 4, row * 3)
            sheet["C5"] = "新しい行の中央"
            sheet["A5"] = "新しい行の先頭"

        base = OOXML.Workbook(book_bytes())
        source = OOXML.Workbook(edit_book(base.data, change))
        selected = {("1", "D" + str(row)) for row in range(1, 301)} | {("1", "C5"), ("1", "A5")}
        result = base.rebuild(source, selected)
        actual = OOXML.Workbook(result)
        for row in range(1, 301):
            self.assertEqual(actual.cells[("1", "D" + str(row))], OOXML.Cell("number", str(row * 3)))
        self.assertEqual(actual.cells[("1", "C1")], OOXML.Cell("string", "残す"))
        root = etree.fromstring(package(result)["xl/worksheets/sheet1.xml"])
        rows = root.find("{" + NS + "}sheetData").findall("{" + NS + "}row")
        self.assertEqual([int(row.get("r")) for row in rows], list(range(1, 301)))
        fifth = [cell.get("r") for cell in rows[4].findall("{" + NS + "}c")]
        self.assertEqual(fifth, ["A5", "C5", "D5"])


if __name__ == "__main__":
    unittest.main()
