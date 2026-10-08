"""Exercise partial Excel staging against real, isolated Git indexes.

These tests select the same projected records that lazygit stages, without
requiring a terminal or altering the developer's Git configuration.
"""

import datetime as dt
import importlib.util
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "excel-stage.py"


def isolated_env(home):
    env = {key: value for key, value in os.environ.items() if not key.startswith("GIT_")}
    env.update(
        HOME=str(home),
        XDG_CONFIG_HOME=str(home / ".config"),
        GIT_CONFIG_NOSYSTEM="1",
        GIT_CONFIG_GLOBAL=os.devnull,
        GIT_TERMINAL_PROMPT="0",
    )
    return env


def save_book(path):
    book = Workbook()
    sales = book.active
    sales.title = "売上"
    sales["A1"] = "元の値"
    sales["B1"] = 10
    sales["C1"] = True
    sales["D1"] = "=B1*2"
    sales["E1"] = "削除する値"
    sales["F1"] = "#DIV/0!"
    sales["A30"] = dt.date(2026, 10, 8)
    sales["A1"].font = Font(bold=True, color="FF0000")
    sales["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    book.create_sheet("備考")["B4"] = "日本語\n改行\tタブ"
    book.save(path)
    book.close()


@unittest.skipUnless(shutil.which("git"), "Git is required for partial staging tests")
class ExcelStageTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # The executable has a hyphen in its filename; import its public session
        # functions rather than replacing them with a fake Git implementation.
        cls.scripts_path = str(SCRIPT.parent)
        sys.path.insert(0, cls.scripts_path)
        spec = importlib.util.spec_from_file_location("excel_stage_integration", SCRIPT)
        cls.stage = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = cls.stage
        spec.loader.exec_module(cls.stage)

    @classmethod
    def tearDownClass(cls):
        sys.path.remove(cls.scripts_path)
        sys.modules.pop("excel_stage_integration", None)

    def setUp(self):
        directory = tempfile.TemporaryDirectory(prefix="excel-stage-git-")
        self.addCleanup(directory.cleanup)
        self.directory = Path(directory.name)
        self.repo = self.directory / "repository"
        self.repo.mkdir()
        self.home = self.directory / "home"
        self.home.mkdir()
        self.env = isolated_env(self.home)
        environment = mock.patch.dict(os.environ, self.env, clear=True)
        environment.start()
        self.addCleanup(environment.stop)
        self.git("init", "--quiet")
        self.git("config", "user.name", "Excel staging test")
        self.git("config", "user.email", "excel-stage@example.invalid")
        self.git("config", "core.autocrlf", "false")
        self.path = self.repo / "日本語 と空白.xlsx"
        save_book(self.path)
        (self.repo / "notes.txt").write_text("original\n", encoding="utf-8")
        self.git("add", "--", self.path.name, "notes.txt")
        self.git("commit", "--quiet", "-m", "Initial workbook")

    def git(self, *args, cwd=None, input=None, env=None, expected_returncode=0):
        result = subprocess.run(
            ["git", "--no-pager", *args],
            cwd=cwd or self.repo,
            env=env or self.env,
            input=input,
            capture_output=True,
            timeout=30,
        )
        self.assertEqual(result.returncode, expected_returncode, result.stderr.decode(errors="replace"))
        return result.stdout

    def edit(self, changes, path=None):
        path = path or self.path
        book = load_workbook(path)
        for (sheet, coordinate), value in changes.items():
            book[sheet][coordinate] = value
        book.save(path)
        book.close()

    def prepare(self):
        session = self.directory / "session"
        state = self.stage.prepare_session(self.repo, self.path.name, session)
        return session, state

    def select(self, session, state, selection):
        """Stage only requested (sheet name, coordinate) projected replacements."""
        picker = session / "picker"
        picker_env = {key: value for key, value in self.env.items() if key != "GIT_INDEX_FILE"}
        for sheet in state["sheets"]:
            filename = sheet["filename"]
            baseline = self.git("show", "HEAD:" + filename, cwd=picker, env=picker_env).decode("utf-8")
            working = (picker / filename).read_text(encoding="utf-8")
            replacements = {
                json.loads(line)[0]: line
                for line in working.splitlines()
                if (sheet["name"], json.loads(line)[0]) in selection
            }
            result = [
                replacements.get(json.loads(line)[0], line)
                for line in baseline.splitlines()
            ]
            (picker / filename).write_text("\n".join(result) + "\n", encoding="utf-8")
            self.git("add", "--", filename, cwd=picker, env=picker_env)

    def index_bytes(self):
        index = self.git("rev-parse", "--git-path", "index").decode().strip()
        return (self.repo / index).read_bytes()

    def staged_bytes(self, path=None):
        return self.git("show", ":" + (path or self.path).name)

    def staged_values(self):
        book = load_workbook(io.BytesIO(self.staged_bytes()))
        self.addCleanup(book.close)
        return book

    def assert_rejected_without_index_change(self, operation):
        before = self.index_bytes()
        with self.assertRaises(ValueError):
            operation()
        self.assertEqual(self.index_bytes(), before)

    def test_multiple_sheets_types_additions_and_clear_preserve_unselected_cells(self):
        self.edit({
            ("売上", "A1"): "変更後",
            ("売上", "B1"): 22.5,
            ("売上", "C1"): False,
            ("売上", "D1"): "=B1*3",
            ("売上", "E1"): None,
            ("売上", "G1"): "新しいセル",
            ("売上", "F1"): "#N/A",
            ("売上", "A30"): dt.date(2026, 10, 9),
            ("備考", "B4"): "別シートの値",
        })
        working_bytes = self.path.read_bytes()
        session, state = self.prepare()
        selection = {
            ("売上", "B1"), ("売上", "C1"), ("売上", "D1"),
            ("売上", "E1"), ("売上", "G1"), ("売上", "F1"), ("売上", "A30"),
            ("備考", "B4"),
        }
        self.select(session, state, selection)
        self.assertEqual(self.stage.apply_session(session), len(selection))
        book = self.staged_values()
        self.assertEqual(book["売上"]["A1"].value, "元の値")
        for coordinate, value in {
            "B1": 22.5, "C1": False, "D1": "=B1*3", "E1": None,
            "G1": "新しいセル", "F1": "#N/A",
        }.items():
            self.assertEqual(book["売上"][coordinate].value, value)
        self.assertEqual(book["備考"]["B4"].value, "別シートの値")
        self.assertEqual(book["売上"]["A30"].value, dt.datetime(2026, 10, 9))
        self.assertEqual(self.path.read_bytes(), working_bytes)

    def test_existing_workbook_stage_and_other_file_stage_survive_apply(self):
        self.edit({("売上", "A1"): "ステージ済み", ("売上", "B1"): 20})
        self.git("add", "--", self.path.name)
        self.edit({("売上", "A1"): "まだ未選択", ("売上", "B1"): 30})
        session, state = self.prepare()
        self.select(session, state, {("売上", "B1")})
        (self.repo / "notes.txt").write_text("staged after session\n", encoding="utf-8")
        self.git("add", "--", "notes.txt")
        notes_before = self.git("show", ":notes.txt")
        self.assertEqual(self.stage.apply_session(session), 1)
        book = self.staged_values()
        self.assertEqual(book["売上"]["A1"].value, "ステージ済み")
        self.assertEqual(book["売上"]["B1"].value, 30)
        self.assertEqual(self.git("show", ":notes.txt"), notes_before)

    def test_japanese_sheet_paths_disable_quoting_only_in_picker(self):
        # delta's OSC filename metadata must contain the same UTF-8 filename
        # lazygit sees; octal-quoted Japanese names make hunk selection a no-op.
        self.git("config", "core.quotePath", "true")
        book = load_workbook(self.path)
        book["売上"].title = "売上 表"
        book.save(self.path)
        book.close()
        self.git("add", "--", self.path.name)
        self.git("commit", "--quiet", "-m", "Japanese sheet name with space")
        self.edit({("売上 表", "A1"): "UTF-8 のハンク選択"})
        session, state = self.prepare()
        sheet = next(sheet for sheet in state["sheets"] if sheet["name"] == "売上 表")
        self.assertIn("売上 表", sheet["filename"])
        self.assertEqual(self.git("config", "--get", "core.quotePath", cwd=session / "picker"), b"false\n")
        self.assertEqual(self.git("config", "--get", "core.quotePath"), b"true\n")
        diff = self.git("diff", "--", sheet["filename"], cwd=session / "picker")
        self.assertIn(sheet["filename"].encode("utf-8"), diff)
        self.assertNotIn(b"\\345", diff)
        self.select(session, state, {("売上 表", "A1")})
        self.assertEqual(self.stage.apply_session(session), 1)
        self.assertEqual(self.staged_values()["売上 表"]["A1"].value, "UTF-8 のハンク選択")
        self.assertEqual(self.git("config", "--get", "core.quotePath"), b"true\n")

    def test_apply_session_cli_stages_selected_values(self):
        self.edit({("売上", "A1"): "CLI から選択"})
        session, state = self.prepare()
        self.select(session, state, {("売上", "A1")})
        result = subprocess.run(
            [sys.executable, str(SCRIPT), "--apply-session", str(session)],
            cwd=self.repo,
            env=self.env,
            capture_output=True,
            timeout=30,
        )
        self.assertEqual(result.returncode, 0, result.stderr.decode(errors="replace"))
        self.assertEqual(self.staged_values()["売上"]["A1"].value, "CLI から選択")

    def test_linked_worktree_updates_its_own_index_and_preserves_main_index(self):
        worktree = self.directory / "別の 作業ツリー"
        self.git("worktree", "add", "--quiet", "-b", "partial-staging", str(worktree))
        main_index = self.index_bytes()
        main_workbook = self.path.read_bytes()
        work_path = worktree / self.path.name
        self.edit({("売上", "A1"): "作業ツリーの変更"}, path=work_path)
        session = self.directory / "linked-session"
        state = self.stage.prepare_session(worktree, work_path.name, session)
        self.assertEqual(Path(state["repo"]), worktree)
        self.assertNotEqual(Path(state["index"]), self.repo / ".git" / "index")
        self.select(session, state, {("売上", "A1")})
        self.assertEqual(self.stage.apply_session(session), 1)
        data = self.git("show", ":" + work_path.name, cwd=worktree)
        book = load_workbook(io.BytesIO(data))
        self.assertEqual(book["売上"]["A1"].value, "作業ツリーの変更")
        book.close()
        self.assertEqual(self.index_bytes(), main_index)
        self.assertEqual(self.path.read_bytes(), main_workbook)

    def test_alternate_index_is_respected_and_isolated_from_picker(self):
        standard_index = self.index_bytes()
        alternate = self.directory / "別の インデックス"
        alternate.write_bytes(standard_index)
        self.edit({("売上", "A1"): "別インデックスの選択"})
        with mock.patch.dict(self.env, {"GIT_INDEX_FILE": str(alternate)}), mock.patch.dict(
            os.environ, {"GIT_INDEX_FILE": str(alternate)}
        ):
            session, state = self.prepare()
            self.assertEqual(Path(state["index"]), alternate)
            self.select(session, state, {("売上", "A1")})
            # Bytes keep LF on Windows, where write_text would store CRLF.
            (self.repo / "notes.txt").write_bytes(b"other alternate stage\n")
            self.git("add", "--", "notes.txt")
            self.assertEqual(self.stage.apply_session(session), 1)
            self.assertEqual(self.staged_values()["売上"]["A1"].value, "別インデックスの選択")
            self.assertEqual(self.git("show", ":notes.txt"), b"other alternate stage\n")
        self.assertEqual(self.index_bytes(), standard_index)

    def test_filename_leading_dash_and_pathspec_characters_are_literal(self):
        filenames = ("-日本語 表.xlsx", "literal[1]*?.xlsx", ":(glob)book.xlsx")
        if os.name == "nt":
            # Windows filenames cannot contain *, ? or :; brackets are still a glob.
            filenames = ("-日本語 表.xlsx", "literal[1].xlsx")
        for number, filename in enumerate(filenames):
            with self.subTest(filename=filename):
                path = self.repo / filename
                save_book(path)
                literal_env = self.env.copy()
                literal_env["GIT_LITERAL_PATHSPECS"] = "1"
                self.git("add", "--", path.name, env=literal_env)
                self.git("commit", "--quiet", "-m", "Literal filename fixture")
                self.edit({("売上", "A1"): filename}, path=path)
                session = self.directory / ("literal-" + str(number))
                state = self.stage.prepare_session(self.repo, filename, session)
                self.select(session, state, {("売上", "A1")})
                self.assertEqual(self.stage.apply_session(session), 1)
                book = load_workbook(io.BytesIO(self.staged_bytes(path)))
                self.assertEqual(book["売上"]["A1"].value, filename)
                book.close()

    @unittest.skipIf(os.name == "nt", "Executable fake tools use POSIX shebangs")
    def test_public_alias_from_subdirectory_cancels_without_altering_index(self):
        setup_spec = importlib.util.spec_from_file_location(
            "excel_stage_setup_integration", SCRIPT.with_name("setup-excel-diff.py")
        )
        setup = importlib.util.module_from_spec(setup_spec)
        setup_spec.loader.exec_module(setup)
        tools = self.directory / "日本語 tools"
        tools.mkdir()
        helper = tools / SCRIPT.name
        shutil.copyfile(SCRIPT, helper)
        shutil.copyfile(SCRIPT.with_name("excel_stage_ooxml.py"), tools / "excel_stage_ooxml.py")
        self.git("config", "alias.excel-stage", setup.stage_alias(Path(sys.executable), helper))
        capture = self.directory / "picker-launch.json"
        fake = (
            f"#!{sys.executable}\n"
            "import json, os, pathlib, sys\n"
            "if '--version' in sys.argv:\n"
            "    print('version=0.66.0' if pathlib.Path(sys.argv[0]).name == 'lazygit' else 'delta 0.20.1')\n"
            "else:\n"
            "    session = pathlib.Path.cwd().parent\n"
            "    state = json.loads((session / 'session.json').read_text(encoding='utf-8'))\n"
            "    config = json.loads((session / 'config.json').read_text(encoding='utf-8'))\n"
            "    data = {'path':state['path'], 'args':sys.argv[1:], 'config':config, 'cwd':str(pathlib.Path.cwd())}\n"
            "    pathlib.Path(os.environ['EXCEL_STAGE_TEST_CAPTURE']).write_text(json.dumps(data), encoding='utf-8')\n"
        )
        for tool in ("lazygit", "delta"):
            executable = tools / tool
            executable.write_text(fake, encoding="utf-8")
            executable.chmod(0o755)
        cli_env = self.env.copy()
        cli_env["PATH"] = str(tools) + os.pathsep + self.env.get("PATH", "")
        cli_env["EXCEL_STAGE_TEST_CAPTURE"] = str(capture)
        folder = self.repo / "subdirectory"
        folder.mkdir()
        path = folder / "-日本語 表.xlsx"
        save_book(path)
        self.git("add", "--", "subdirectory/" + path.name)
        self.git("commit", "--quiet", "-m", "Workbook below caller directory")
        self.edit({("売上", "A1"): "キャンセル"}, path=path)
        before = self.index_bytes()
        stdout = self.git("excel-stage", "--", path.name, cwd=folder, env=cli_env)
        self.assertIn("選択を取り消しました", stdout.decode("utf-8"))
        launch = json.loads(capture.read_text(encoding="utf-8"))
        self.assertEqual(launch["path"], "subdirectory/" + path.name)
        self.assertIn("--use-config-dir", launch["args"])
        self.assertIn("--use-config-file", launch["args"])
        self.assertTrue(launch["config"]["gui"]["useHunkModeInDiffView"])
        self.assertIn("--side-by-side", launch["config"]["git"]["diffRenderers"][0]["command"])
        self.assertFalse(Path(launch["cwd"]).exists(), "cancelled session must be cleaned up")
        self.assertEqual(self.index_bytes(), before)

    def test_cancel_and_apply_with_no_selection_preserve_index_exactly(self):
        self.edit({("売上", "A1"): "選ばない"})
        before = self.index_bytes()
        session, _ = self.prepare()
        # Closing the picker performs no apply operation.
        self.assertEqual(self.index_bytes(), before)
        self.assertEqual(self.stage.apply_session(session), 0)
        self.assertEqual(self.index_bytes(), before)

    def test_target_index_change_rejects_and_preserves_newer_stage(self):
        self.edit({("売上", "A1"): "変更後"})
        session, state = self.prepare()
        self.select(session, state, {("売上", "A1")})
        self.git("add", "--", self.path.name)
        self.assert_rejected_without_index_change(lambda: self.stage.apply_session(session))

    def test_working_file_change_rejects_and_preserves_index(self):
        self.edit({("売上", "A1"): "変更後"})
        session, state = self.prepare()
        self.select(session, state, {("売上", "A1")})
        self.edit({("売上", "B1"): 999})
        self.assert_rejected_without_index_change(lambda: self.stage.apply_session(session))

    def test_working_file_deletion_rejects_and_preserves_index(self):
        self.edit({("売上", "A1"): "変更後"})
        session, state = self.prepare()
        self.select(session, state, {("売上", "A1")})
        self.path.unlink()
        self.assert_rejected_without_index_change(lambda: self.stage.apply_session(session))

    def test_existing_index_lock_rejects_without_removing_external_lock(self):
        self.edit({("売上", "A1"): "変更後"})
        session, state = self.prepare()
        self.select(session, state, {("売上", "A1")})
        lock = self.repo / ".git" / "index.lock"
        lock.write_bytes(b"another Git process")
        self.assert_rejected_without_index_change(lambda: self.stage.apply_session(session))
        self.assertEqual(lock.read_bytes(), b"another Git process")

    def test_incomplete_replacement_halves_are_rejected(self):
        for malformed in ("missing", "duplicate"):
            with self.subTest(malformed=malformed):
                self.edit({("売上", "A1"): "変更後"})
                session = self.directory / malformed
                state = self.stage.prepare_session(self.repo, self.path.name, session)
                sheet = next(sheet for sheet in state["sheets"] if sheet["name"] == "売上")
                picker = session / "picker"
                path = picker / sheet["filename"]
                baseline = self.git("show", "HEAD:" + sheet["filename"], cwd=picker).decode()
                working = path.read_text(encoding="utf-8")
                if malformed == "missing":
                    records = [line for line in baseline.splitlines() if json.loads(line)[0] != "A1"]
                else:
                    replacement = next(line for line in working.splitlines() if json.loads(line)[0] == "A1")
                    records = baseline.splitlines() + [replacement]
                path.write_text("\n".join(records) + "\n", encoding="utf-8")
                self.git("add", "--", sheet["filename"], cwd=picker)
                self.assert_rejected_without_index_change(lambda: self.stage.apply_session(session))

    def test_invented_record_value_is_rejected(self):
        self.edit({("売上", "A1"): "変更後"})
        session, state = self.prepare()
        sheet = next(sheet for sheet in state["sheets"] if sheet["name"] == "売上")
        picker = session / "picker"
        path = picker / sheet["filename"]
        records = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]
        next(record for record in records if record[0] == "A1")[2] = "not either original value"
        path.write_text("".join(json.dumps(record, ensure_ascii=False) + "\n" for record in records), encoding="utf-8")
        self.git("add", "--", sheet["filename"], cwd=picker)
        self.assert_rejected_without_index_change(lambda: self.stage.apply_session(session))

    def test_unsupported_and_new_files_require_whole_file_staging(self):
        for extension in (".xls", ".xlt", ".txt"):
            with self.subTest(extension=extension):
                path = self.repo / ("old-format" + extension)
                path.write_bytes(self.path.read_bytes())
                self.git("add", "--", path.name)
                self.git("commit", "--quiet", "-m", "Unsupported format fixture")
                self.assert_rejected_without_index_change(
                    lambda: self.stage.prepare_session(self.repo, path.name, self.directory / extension[1:])
                )
        new = self.repo / "new.xlsx"
        save_book(new)
        for staged in (False, True):
            with self.subTest(staged=staged):
                if staged:
                    self.git("add", "--", new.name)
                self.assert_rejected_without_index_change(
                    lambda: self.stage.prepare_session(self.repo, new.name, self.directory / ("new-" + str(staged)))
                )

    def test_supported_ooxml_extensions_accept_partial_staging(self):
        for extension in (".xlsx", ".xlsm", ".xltx", ".xltm"):
            with self.subTest(extension=extension):
                path = self.repo / ("形式" + extension)
                save_book(path)
                self.git("add", "--", path.name)
                self.git("commit", "--quiet", "-m", "OOXML extension fixture")
                self.edit({("売上", "B1"): 25}, path=path)
                session = self.directory / ("format-" + extension[1:])
                state = self.stage.prepare_session(self.repo, path.name, session)
                self.select(session, state, {("売上", "B1")})
                self.assertEqual(self.stage.apply_session(session), 1)
                book = load_workbook(io.BytesIO(self.staged_bytes(path)))
                self.assertEqual(book["売上"]["B1"].value, 25)
                self.assertEqual(book["売上"]["A1"].value, "元の値")
                book.close()

    def test_deleted_and_corrupt_workbooks_are_rejected(self):
        self.path.unlink()
        self.assert_rejected_without_index_change(self.prepare)
        self.path.write_bytes(b"PK\x03\x04broken workbook")
        self.assert_rejected_without_index_change(self.prepare)

    def test_corrupt_index_workbook_is_rejected(self):
        self.path.write_bytes(b"not a workbook")
        self.git("add", "--", self.path.name)
        save_book(self.path)
        self.assert_rejected_without_index_change(self.prepare)

    def test_unmerged_workbook_is_rejected(self):
        oid = self.git("rev-parse", ":" + self.path.name).decode().strip()
        self.git("update-index", "--force-remove", "--", self.path.name)
        entries = "".join(f"100644 {oid} {stage}\t{self.path.name}\n" for stage in (1, 2, 3))
        self.git("update-index", "--index-info", input=entries.encode())
        self.assert_rejected_without_index_change(self.prepare)

    def test_sheet_rename_add_delete_and_order_change_are_rejected(self):
        baseline = self.path.read_bytes()
        for change in ("rename", "add", "delete", "reorder"):
            with self.subTest(change=change):
                self.path.write_bytes(baseline)
                book = load_workbook(self.path)
                if change == "rename":
                    book["売上"].title = "別名"
                elif change == "add":
                    book.create_sheet("追加")
                elif change == "delete":
                    del book["備考"]
                else:
                    book.move_sheet("備考", offset=-1)
                book.save(self.path)
                book.close()
                self.assert_rejected_without_index_change(
                    lambda: self.stage.prepare_session(self.repo, self.path.name, self.directory / change)
                )

    def test_formatting_only_changes_cannot_be_selected(self):
        baseline = self.staged_bytes()
        book = load_workbook(self.path)
        book["売上"]["A1"].font = Font(italic=True, color="0000FF")
        book["売上"].column_dimensions["A"].width = 70
        book.save(self.path)
        book.close()
        self.assertNotEqual(self.path.read_bytes(), baseline)
        session, state = self.prepare()
        self.assertEqual(self.git("diff", "--", cwd=session / "picker"), b"")
        self.assertEqual(self.stage.apply_session(session), 0)
        self.assertEqual(self.staged_bytes(), baseline)

    def test_index_style_is_preserved_when_selected_cell_value_changes(self):
        self.edit({("売上", "A1"): "新しい文字列"})
        book = load_workbook(self.path)
        book["売上"]["A1"].font = Font(italic=True, color="0000FF")
        book.save(self.path)
        book.close()
        session, state = self.prepare()
        self.select(session, state, {("売上", "A1")})
        self.assertEqual(self.stage.apply_session(session), 1)
        staged = self.staged_values()["売上"]["A1"]
        self.assertEqual(staged.value, "新しい文字列")
        self.assertTrue(staged.font.bold)
        self.assertFalse(staged.font.italic)
        self.assertEqual(staged.font.color.rgb, "00FF0000")
        self.assertEqual(staged.fill.fgColor.rgb, "00FFFF00")


if __name__ == "__main__":
    unittest.main()
