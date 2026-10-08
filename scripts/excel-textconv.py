#!/usr/bin/env python3
"""Render Excel cells as stable UTF-8 text for Git's textconv diff driver."""

import argparse
import datetime as dt
import json
import sys
from pathlib import Path


OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def quoted(value):
    # Escape tabs/newlines so a cell always occupies exactly one diff line.
    return json.dumps(value, ensure_ascii=False, allow_nan=False)


def value_text(value):
    if isinstance(value, dt.datetime):
        return "datetime: " + value.isoformat()
    if isinstance(value, dt.date):
        return "date: " + value.isoformat()
    if isinstance(value, dt.time):
        return "time: " + value.isoformat()
    if isinstance(value, dt.timedelta):
        return "duration-seconds: " + quoted(value.total_seconds())
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return quoted(value)


def render_ooxml(source):
    from openpyxl import load_workbook
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

    # Git's temporary blobs have no Excel extension. Passing an open stream also
    # lets openpyxl read .xlsm/.xltx/.xltm without relying on the filename.
    book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    try:
        for sheet in book:
            print("Sheet: " + quoted(sheet.title))
            # Read actual rows, including files with incorrect dimension metadata.
            sheet.reset_dimensions()
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    value = cell.value
                    if isinstance(value, ArrayFormula):
                        text = "array-formula: " + quoted(
                            {"text": value.text, "ref": value.ref}
                        )
                    elif isinstance(value, DataTableFormula):
                        text = "data-table-formula: " + quoted(dict(value))
                    elif cell.data_type == "f":
                        text = "formula: " + quoted(value)
                    elif cell.data_type == "e":
                        text = "error: " + quoted(value)
                    else:
                        text = value_text(value)
                    print(cell.coordinate + "\t" + text)
    finally:
        book.close()


def render_xls(source):
    import xlrd
    from openpyxl.utils.cell import get_column_letter

    # xlrd exposes cached formula results, not the formula expressions in .xls.
    book = xlrd.open_workbook(file_contents=source.read(), on_demand=True)
    try:
        for sheet in book.sheets():
            print("Sheet: " + quoted(sheet.name))
            for row_index in range(sheet.nrows):
                for col_index, cell in enumerate(sheet.row(row_index)):
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                        if 0 <= cell.value < 1:
                            value = value.time()
                        text = value_text(value)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        text = quoted(bool(value))
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        text = "error: " + quoted(xlrd.error_text_from_code[value])
                    else:
                        text = value_text(value)
                    coordinate = get_column_letter(col_index + 1) + str(row_index + 1)
                    print(coordinate + "\t" + text)
    finally:
        book.release_resources()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path, help="Excel file or Git temporary blob")
    args = parser.parse_args()
    # Git consumes UTF-8 even on Windows with a non-UTF-8 console code page.
    sys.stdout.reconfigure(encoding="utf-8", newline="\n")
    sys.stderr.reconfigure(encoding="utf-8")
    try:
        with args.file.open("rb") as source:
            signature = source.read(8)
            source.seek(0)
            if signature.startswith(b"PK"):
                render_ooxml(source)
            elif signature == OLE_SIGNATURE:
                render_xls(source)
            else:
                raise ValueError("対応する Excel ファイルではありません (xlsx/xlsm/xltx/xltm/xls/xlt)")
    except (Exception, KeyboardInterrupt) as error:
        print("Excel の差分変換に失敗しました: " + str(error), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
